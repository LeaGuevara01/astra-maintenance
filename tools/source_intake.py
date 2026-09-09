"""Local-only source intake. Originals are never modified or uploaded."""
import argparse, hashlib, json, re, sys, zipfile
from pathlib import Path
from collections import Counter
from xml.etree import ElementTree as ET
import pymupdf
import openpyxl
from PIL import Image

def extract(path, extension):
    if extension==".pdf":
        doc=pymupdf.open(path)
        chunks=[]
        scanned=[]
        for n,page in enumerate(doc):
            text=page.get_text("text")
            chunks.append(f"\n--- PAGE {n+1} ---\n{text}")
            if len(text.strip())<20: scanned.append(n+1)
        return "".join(chunks),{"pages":len(doc),"pagesNeedingOCR":scanned}
    if extension==".docx":
        with zipfile.ZipFile(path) as z:
            chunks=[]
            for name in z.namelist():
                if re.match(r"word/(document|header\d+|footer\d+)\.xml$",name):
                    tree=ET.fromstring(z.read(name))
                    chunks.append("\n".join("".join(p.itertext()) for p in tree.findall(".//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")))
            return "\n".join(chunks),{}
    if extension in (".xlsx",".xlsm"):
        wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
        chunks=[]
        for sh in wb:
            chunks.append(f"\n--- SHEET {sh.title} ---")
            for row in sh.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    chunks.append("\t".join(str(v) if v is not None else "" for v in row))
        sheets=wb.sheetnames;wb.close()
        return "\n".join(chunks),{"sheets":sheets}
    if extension in (".csv",".tsv",".txt",".md",".json"):
        data=path.read_bytes()
        for encoding in ("utf-8-sig","utf-16","cp1252"):
            try:return data.decode(encoding),{}
            except UnicodeError:pass
        return data.decode("utf-8",errors="replace"),{}
    if extension in (".jpg",".jpeg",".png",".webp",".tif",".tiff"):
        with Image.open(path) as im:return "",{"width":im.width,"height":im.height,"format":im.format}
    return "",{}

def main():
    a=argparse.ArgumentParser();a.add_argument("--inventory",required=True);a.add_argument("--output",required=True)
    args=a.parse_args();out=Path(args.output);(out/"text").mkdir(parents=True,exist_ok=True)
    raw=json.loads(Path(args.inventory).read_text(encoding="utf-8-sig"))
    selected=[]
    keywords=r"astra|manten|tractor|taller|repuesto|spare|stock|agric|maquina|tarjeta|puma|lubric|filtro|manual|gestion|cat.logo|deutz|new.holland|john.deere|\\TÉCNICAS\\|\\DOCUMENTACIÓN\\|\\CATÁLOGOS Y MANUALES\\|\\Pictures\\(caja|maquinas|camiones|Piezas|burros|bisagra|Diagramas)"
    extensions={".pdf",".docx",".xlsx",".xlsm",".md",".txt",".csv",".tsv",".json",".jpg",".jpeg",".png",".webp",".tif",".tiff"}
    for item in raw:
        path=item["path"]
        if re.search(r"\\Image-Line\\|\\node_modules\\|\\.git\\|\\Codex\\",path,re.I):continue
        if item["ext"].lower() in extensions and re.search(keywords,path,re.I):selected.append(item)
    index=[];seen={}
    for i,item in enumerate(selected):
        path=Path(item["path"]);entry=dict(item)
        try:
            with path.open("rb") as handle:
                digest=hashlib.file_digest(handle,"sha256").hexdigest()
            entry.update(id="SRC-"+digest[:12],sha256=digest,reviewStatus="UNREVIEWED")
            name=path.name
            if re.search(r"factura|presupuesto|remito|pago|precio",name,re.I):kind="COMMERCIAL_HISTORY"
            elif re.search(r"manual|cat.logo",name,re.I):kind="MANUAL_OR_CATALOG"
            elif re.search(r"tarjeta|matriz|protocolo|planilla|astra",name,re.I):kind="INTERNAL_WORKFLOW"
            else:kind="TECHNICAL_REFERENCE"
            entry["kind"]=kind
            if digest in seen:
                entry.update(extractionStatus="DUPLICATE",duplicateOf=seen[digest])
            else:
                seen[digest]=entry["id"]
                text,extra=extract(path,path.suffix.lower())
                entry.update(extra)
                status="TEXT_EXTRACTED" if text.strip() else "VISUAL_REVIEW_REQUIRED"
                if path.suffix.lower()==".pdf" and len(re.sub(r"--- PAGE \d+ ---","",text).strip())<20:status="OCR_REQUIRED"
                if text.strip():
                    dest=out/"text"/(entry["id"]+".txt");dest.write_text(text,encoding="utf-8")
                    entry["textFile"]="text/"+dest.name
                    entry["textCharacters"]=len(text)
                entry["extractionStatus"]=status
        except Exception as exc:
            entry.update(extractionStatus="ERROR",error=str(exc)[:250])
        index.append(entry)
        if i%50==0:print(f"Indexed {i+1}/{len(selected)}",flush=True)
    (out/"index.json").write_text(json.dumps(index,ensure_ascii=False,indent=2),encoding="utf-8")
    summary={"inventoriedFiles":len(raw),"selected":len(index),"uniqueHashes":len(seen),"status":dict(Counter(x["extractionStatus"] for x in index)),"kind":dict(Counter(x.get("kind","unknown") for x in index)),"note":"Extraction is not technical validation. Source originals stay local; OCR/visual review is separate."}
    (out/"summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(summary,ensure_ascii=False))
if __name__=="__main__":main()
